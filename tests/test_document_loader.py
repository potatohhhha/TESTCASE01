from pathlib import Path

import pytest
from docx import Document as DocxDocument

from autotest.document_loader import DocumentLoader


def test_reads_markdown(workdir: Path):
    path = workdir / "requirements.md"
    path.write_text("# REQ001\nDo the thing", encoding="utf-8")
    result = DocumentLoader().load(path)
    assert result.file_type == "markdown"
    assert "REQ001" in result.text


def test_rejects_empty_and_unsupported(workdir: Path):
    empty = workdir / "empty.txt"
    empty.write_text("  ", encoding="utf-8")
    with pytest.raises(ValueError, match="empty"):
        DocumentLoader().load(empty)
    unsupported = workdir / "requirements.pdf"
    unsupported.write_bytes(b"x")
    with pytest.raises(ValueError, match="Unsupported"):
        DocumentLoader().load(unsupported)


def test_reads_txt(workdir: Path):
    path = workdir / "requirements.txt"
    path.write_text("REQ-TXT: plain text requirement", encoding="utf-8")
    result = DocumentLoader().load(path)
    assert result.file_type == "txt"
    assert "REQ-TXT" in result.text


def test_reads_docx_paragraphs_and_tables(workdir: Path):
    path = workdir / "requirements.docx"
    source = DocxDocument()
    source.add_paragraph("REQ-DOCX: paragraph requirement")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "REQ-TABLE"
    table.cell(0, 1).text = "table requirement"
    source.save(path)
    result = DocumentLoader().load(path)
    assert result.file_type == "docx"
    assert "REQ-DOCX" in result.text
    assert "REQ-TABLE\ttable requirement" in result.text

