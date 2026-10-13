"""文件功能：读取 TXT、Markdown 和 DOCX 格式的需求文档。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class Document:
    text: str
    file_path: Path
    file_type: str


class DocumentLoader:
    """读取需求文档，但不对文档内容作语义解释。"""

    TEXT_TYPES = {".txt": "txt", ".md": "markdown", ".markdown": "markdown"}

    def load(self, file_path: str | Path) -> Document:
        path = Path(file_path).resolve()
        if not path.is_file():
            raise FileNotFoundError(f"Requirement document does not exist: {path}")
        suffix = path.suffix.lower()
        if suffix in self.TEXT_TYPES:
            text = path.read_text(encoding="utf-8-sig")
            kind = self.TEXT_TYPES[suffix]
        elif suffix == ".docx":
            from docx import Document as DocxDocument
            doc = DocxDocument(path)
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
            text, kind = "\n".join(parts), "docx"
        else:
            raise ValueError(f"Unsupported requirement document format: {suffix or '(none)'}")
        if not text.strip():
            raise ValueError(f"Requirement document is empty: {path}")
        return Document(text=text, file_path=path, file_type=kind)

